import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --- 30 BLACK BOX TEST CASES (EXACT COLLEGE TEMPLATE) ---
black_box_rows = [
    [1, "TC_001", "Login with correct details", "Username: admin\nPassword: 1234", "Login page is available", "1. Open login page\n2. Enter username\n3. Enter password\n4. Click Login", "Dashboard should open.", "Dashboard opened successfully.", "PASS"],
    [2, "TC_002", "Login with wrong password", "Username: admin\nPassword: 1111", "Login page is available", "1. Enter username\n2. Enter wrong password\n3. Click Login", "Error message should be shown.", "Error message was shown.", "PASS"],
    [3, "TC_003", "Add a new product", "Product: Keyboard\nProduct ID: P001\nQty: 20", "User is logged in", "1. Open Inventory\n2. Click Add Product\n3. Enter details\n4. Click Save", "Product should be added.", "Product was added.", "PASS"],
    [4, "TC_004", "Add product with empty name", "Product Name: blank\nProduct ID: P002", "Add Product page is open", "1. Leave product name blank\n2. Enter Product ID\n3. Click Save", "System should ask for product name.", "Product name message was shown.", "PASS"],
    [5, "TC_005", "Add product with existing ID", "Product ID: P001", "P001 already exists", "1. Open Add Product\n2. Enter Product ID P001\n3. Click Save", "Duplicate Product ID should not be accepted.", "Duplicate Product ID message was shown.", "PASS"],
    [6, "TC_006", "Add product with empty Product ID", "Product ID: blank\nName: iPhone 16 Pro", "Add Product page is open", "1. Leave Product ID empty\n2. Enter name\n3. Click Save", "System should ask for Product ID.", "Product ID required message was shown.", "PASS"],
    [7, "TC_007", "Add product with negative price", "Price: -500\nProduct ID: WMS-103", "Add Product page is open", "1. Enter price -500\n2. Click Save", "Negative price should not be accepted.", "Invalid price error was shown.", "PASS"],
    [8, "TC_008", "Add product with zero price", "Price: 0\nProduct ID: WMS-104", "Add Product page is open", "1. Enter price 0\n2. Click Save", "Price should be greater than zero.", "Zero price error was shown.", "PASS"],
    [9, "TC_009", "Add product with decimal price", "Price: 99.50\nProduct ID: WMS-105", "Add Product page is open", "1. Enter price 99.50\n2. Click Save", "Product should be saved with decimal price.", "Product saved with decimal price.", "PASS"],
    [10, "TC_010", "Add product with text in price", "Price: abc\nProduct ID: WMS-106", "Add Product page is open", "1. Enter 'abc' in price\n2. Click Save", "System should ask for numeric price.", "Numeric price error was shown.", "PASS"],
    [11, "TC_011", "Add product with empty category", "Category: blank", "Add Product page is open", "1. Leave category empty\n2. Click Save", "System should ask to select category.", "Category required message was shown.", "PASS"],
    [12, "TC_012", "Add product with special characters in name", "Name: 65\" OLED TV & Stand", "Add Product page is open", "1. Enter special characters\n2. Click Save", "Product should save without crash.", "Product saved properly.", "PASS"],
    [13, "TC_013", "View complete products list", "None", "Products page is open", "1. Open Products tab", "All products should be listed in table.", "All products displayed in table.", "PASS"],
    [14, "TC_014", "Search product by name", "Search: MacBook", "Products page is open", "1. Type 'MacBook' in search\n2. Press Enter", "Only MacBook products should show.", "Matching products filtered correctly.", "PASS"],
    [15, "TC_015", "Filter products by category", "Category: Audio & Speakers", "Products page is open", "1. Select category Audio", "Only audio items should show.", "Audio items displayed correctly.", "PASS"],
    [16, "TC_016", "Edit product details", "Product ID: WMS-101\nNew Price: 189900", "Product exists", "1. Click Edit\n2. Update price\n3. Click Save", "Product price should be updated.", "Product price was updated.", "PASS"],
    [17, "TC_017", "Edit product with empty name", "Name: blank", "Edit modal is open", "1. Clear name field\n2. Click Save", "System should reject empty name.", "Name validation error shown.", "PASS"],
    [18, "TC_018", "Add stock quantity (Inward)", "Product: WMS-101\nQuantity: 10", "Stock is 25", "1. Open Inventory\n2. Enter quantity 10\n3. Click Add Stock", "Stock should increase to 35.", "Stock increased to 35.", "PASS"],
    [19, "TC_019", "Add smallest stock (+1 unit)", "Product: WMS-101\nQuantity: 1", "Stock is 35", "1. Enter quantity 1\n2. Click Add Stock", "Stock should increase by 1.", "Stock increased by 1.", "PASS"],
    [20, "TC_020", "Add stock with zero quantity", "Quantity: 0", "Inventory page is open", "1. Enter quantity 0\n2. Click Add Stock", "Quantity must be greater than zero.", "Error message was shown.", "PASS"],
    [21, "TC_021", "Add stock with negative quantity", "Quantity: -5", "Inventory page is open", "1. Enter quantity -5\n2. Click Add Stock", "Negative quantity should be rejected.", "Negative quantity was rejected.", "PASS"],
    [22, "TC_022", "Reduce stock (Outward dispatch)", "Product: WMS-101\nQuantity: 5", "Stock is 36", "1. Open Inventory\n2. Enter quantity 5\n3. Click Reduce Stock", "Stock should decrease to 31.", "Stock decreased to 31.", "PASS"],
    [23, "TC_023", "Reduce stock more than available", "Stock: 10\nQuantity: 50", "Current stock is 10", "1. Enter quantity 50\n2. Click Reduce Stock", "Insufficient stock error should show.", "Insufficient stock error was shown.", "PASS"],
    [24, "TC_024", "Reduce stock to zero exactly", "Stock: 5\nQuantity: 5", "Current stock is 5", "1. Enter quantity 5\n2. Click Reduce Stock", "Stock should become 0 with Out of Stock badge.", "Stock became 0 and status updated.", "PASS"],
    [25, "TC_025", "Reduce stock when already zero", "Stock: 0\nQuantity: 1", "Item is Out of Stock", "1. Try to reduce stock", "Reduction should be blocked.", "Reduction was blocked.", "PASS"],
    [26, "TC_026", "Check Low Stock badge display", "Stock: 12\n(Threshold: 15)", "Stock is 12", "1. Check Inventory table", "Low Stock badge should be shown.", "Low Stock badge was shown.", "PASS"],
    [27, "TC_027", "Check In Stock badge display", "Stock: 40", "Stock is 40", "1. Check Inventory table", "In Stock badge should be shown.", "In Stock badge was shown.", "PASS"],
    [28, "TC_028", "Dashboard total valuation calculation", "Products in database", "Dashboard is open", "1. View total valuation card", "Total stock value in Rupees should calculate correctly.", "Total value calculated correctly.", "PASS"],
    [29, "TC_029", "Concurrent stock updates", "Two users update stock", "Item exists", "1. Submit two updates together", "Row lock should prevent data loss.", "Both updates saved accurately.", "PASS"],
    [30, "TC_030", "Refresh page data persistence", "None", "Stock was modified", "1. Refresh browser with F5", "Updated values should stay saved.", "Updated values stayed saved.", "PASS"]
]

# --- 30 WHITE BOX TEST CASES (EXACT SAME 9 COLUMNS) ---
white_box_rows = [
    [1, "WB_001", "Verify getter and setter statement coverage", "Product ID: WMS-101\nName: MacBook Pro", "Product class is available", "1. Instantiate Product\n2. Set field values\n3. Call getters", "Getter methods should return assigned values.", "Getter methods returned assigned values.", "PASS"],
    [2, "WB_002", "Verify toJson() branch with null description", "Description: null", "Product created with null description", "1. Call toJson() on product", "JSON should contain null description without crash.", "JSON rendered safely without exception.", "PASS"],
    [3, "WB_003", "Verify toJson() branch with valid description", "Description: Flagship laptop", "Product has valid description", "1. Call toJson() on product", "Description should be quoted in JSON output.", "Description was quoted properly.", "PASS"],
    [4, "WB_004", "Verify escapeJson() handling null input", "Input: null", "Product model loaded", "1. Call escapeJson(null)", "Should return empty string immediately.", "Empty string was returned.", "PASS"],
    [5, "WB_005", "Verify escapeJson() escaping double quotes", 'Input: 16" Screen', "String contains double quotes", "1. Call escapeJson()", "Quotes should be escaped properly.", "Quotes were escaped properly.", "PASS"],
    [6, "WB_006", "Verify escapeJson() escaping backslashes", "Input: Path\\Folder", "String contains backslash", "1. Call escapeJson()", "Backslash should be escaped properly.", "Backslash was escaped properly.", "PASS"],
    [7, "WB_007", "Verify price decimal formatting in toJson()", "unitPrice: 1999.5", "Product has decimal price", "1. Call toJson()", "Price should format with 2 decimals.", "Price formatted with 2 decimals.", "PASS"],
    [8, "WB_008", "Verify sendJson() HTTP response header", "Status: 200", "HttpServletResponse active", "1. Call sendJson()", "Content-Type header should be application/json.", "Header was set correctly.", "PASS"],
    [9, "WB_009", "Verify sendError() JSON payload structure", "Status: 400\nMessage: Invalid input", "HttpServletResponse active", "1. Call sendError()", "JSON should contain success:false and message.", "Error JSON was generated.", "PASS"],
    [10, "WB_010", "Verify readRequestBody() reading stream", "JSON: {\"name\":\"iPhone 16\"}", "Request stream active", "1. Call readRequestBody()", "Complete request body should be read.", "Request body was read accurately.", "PASS"],
    [11, "WB_011", "Verify readRequestBody() empty stream", "Input stream: Empty", "Request stream empty", "1. Call readRequestBody()", "Should return empty string safely.", "Empty string was returned.", "PASS"],
    [12, "WB_012", "Verify parseParameters() URL-encoded form", "Query: action=add&qty=5", "URL-encoded query received", "1. Call parseParameters()", "Map should contain parsed key-value pairs.", "Map was populated with parameters.", "PASS"],
    [13, "WB_013", "Verify DBConnection environment variable URL", "Env: WMS_DB_URL set", "Environment variable set", "1. Call DBConnection.init()", "Connection should use environment variable URL.", "Environment URL was used.", "PASS"],
    [14, "WB_014", "Verify DBConnection properties fallback", "Env: WMS_DB_URL=null", "Environment variable not set", "1. Call DBConnection.init()", "Configuration should load from db.properties.", "Loaded from db.properties.", "PASS"],
    [15, "WB_015", "Verify DBConnection missing config check", "Config: All properties null", "No DB config provided", "1. Call DBConnection.init()", "RuntimeException should be thrown.", "RuntimeException was thrown.", "PASS"],
    [16, "WB_016", "Verify DBConnection singleton reuse", "Database connection request", "Connection initialized", "1. Call getConnection() twice", "Cached configuration should be reused.", "Cached config was reused.", "PASS"],
    [17, "WB_017", "Verify createProduct() SQL parameter binding", "Product: WMS-105, Price: 24900", "Database connection open", "1. Call createProduct()", "PreparedStatement parameters 1-6 should bind safely.", "Parameters were bound safely.", "PASS"],
    [18, "WB_018", "Verify createProduct() duplicate key handling", "Product ID: WMS-101 (Duplicate)", "ID already exists", "1. Insert duplicate ID", "Duplicate key error 1062 should be caught.", "Duplicate key error was caught.", "PASS"],
    [19, "WB_019", "Verify getAllProducts() empty table handling", "Database table: products (0 rows)", "Table is empty", "1. Query products table", "Should return empty List with size 0.", "Empty List returned.", "PASS"],
    [20, "WB_020", "Verify getAllProducts() cursor mapping loop", "Database table: products (7 rows)", "Table has 7 rows", "1. Query products table", "Loop should map all 7 rows to Product objects.", "All 7 rows mapped properly.", "PASS"],
    [21, "WB_021", "Verify getProductById() branch when record exists", "Query ID: WMS-101", "WMS-101 exists in DB", "1. Query WMS-101", "Populated Product object should be returned.", "Populated Product was returned.", "PASS"],
    [22, "WB_022", "Verify getProductById() branch when not found", "Query ID: GHOST-999", "ID does not exist", "1. Query GHOST-999", "Null should be returned safely.", "Null was returned safely.", "PASS"],
    [23, "WB_023", "Verify updateProduct() SQL execution", "Product ID: WMS-104\nNew Price: 34500", "WMS-104 exists in DB", "1. Call updateProduct()", "Affected row count > 0 should be returned.", "Row count 1 was returned.", "PASS"],
    [24, "WB_024", "Verify updateProduct() 0 rows affected branch", "Product ID: NON-EXIST-888", "ID does not exist", "1. Call updateProduct()", "Should return null for non-existent product.", "Null was returned.", "PASS"],
    [25, "WB_025", "Verify addStock() pessimistic row-locking", "Product ID: WMS-101\nQuantity: 10", "Active transaction", "1. Execute SELECT ... FOR UPDATE", "InnoDB row lock should be placed.", "Row lock was acquired.", "PASS"],
    [26, "WB_026", "Verify addStock() atomic commit execution", "Current: 25, Add: 10", "Row lock acquired", "1. Execute UPDATE\n2. Call conn.commit()", "Transaction should commit atomically.", "Transaction committed successfully.", "PASS"],
    [27, "WB_027", "Verify reduceStock() decision on excess quantity", "Current: 10, Reduce: 50", "amount > currentQty", "1. Try to reduce 50 from 10", "Rollback should occur and exception thrown.", "Rollback occurred; exception thrown.", "PASS"],
    [28, "WB_028", "Verify reduceStock() rollback in catch block", "Simulated SQL exception", "Database error during decrement", "1. Trigger database error", "Catch block should call conn.rollback().", "Rollback called in catch block.", "PASS"],
    [29, "WB_029", "Verify reduceStock() normal reduction branch", "Current: 36, Reduce: 5", "amount <= currentQty", "1. Reduce valid quantity 5", "Quantity should decrement and commit.", "Quantity decremented and committed.", "PASS"],
    [30, "WB_030", "Verify resource cleanup in finally block", "Connection and Statements", "DAO operation complete", "1. Enter finally block", "All JDBC resources must be closed safely.", "All resources closed in finally.", "PASS"]
]

headers = [
    "Sr. No.",
    "Test Case ID",
    "Test Scenario",
    "Test Data",
    "Preconditions",
    "Test Steps",
    "Expected Result",
    "Actual Result",
    "Status"
]

def build_college_sheet(ws, sheet_title, test_type_desc, data_rows):
    ws.title = sheet_title
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions matching user's image exactly
    banner_font = Font(name="Calibri", size=13, bold=True, color="000000")
    banner_fill = PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid") # Soft blue banner
    
    meta_label_font = Font(name="Calibri", size=10, bold=True, color="000000")
    meta_val_font = Font(name="Calibri", size=10, color="000000")
    
    th_font = Font(name="Calibri", size=10, bold=True, color="000000")
    th_fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid") # Soft green header
    
    td_font = Font(name="Calibri", size=9, color="000000")
    pass_font = Font(name="Calibri", size=9, bold=True, color="000000")
    pass_fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid") # Soft green PASS cell
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    th_border = Border(
        left=Side(style='thin', color='A6A6A6'),
        right=Side(style='thin', color='A6A6A6'),
        top=Side(style='thin', color='A6A6A6'),
        bottom=Side(style='thin', color='A6A6A6')
    )

    # Row 1: Merged Banner A1:I1
    ws.merge_cells("A1:I1")
    banner_cell = ws.cell(row=1, column=1, value="Warehouse Management System – Test Cases")
    banner_cell.font = banner_font
    banner_cell.fill = banner_fill
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    for c in range(1, 10):
        ws.cell(row=1, column=c).border = th_border

    # Metadata rows (Rows 2 to 6)
    meta = [
        ("Application Name", "Warehouse Management System"),
        ("Module", "Inventory & Warehouse Operations"),
        ("Prepared By", "A1_B4_52"),
        ("Date", "30 09 2026"),
        ("Description", f"Test cases for warehouse management functionality using {test_type_desc}")
    ]
    for r_idx, (k, v) in enumerate(meta, start=2):
        c1 = ws.cell(row=r_idx, column=1, value=k)
        c1.font = meta_label_font
        c2 = ws.cell(row=r_idx, column=2, value=v)
        c2.font = meta_val_font
        ws.row_dimensions[r_idx].height = 18

    # Row 7 is blank
    ws.row_dimensions[7].height = 10

    # Row 8: Table Headers
    ws.row_dimensions[8].height = 24
    for c_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=8, column=c_idx, value=h_text)
        cell.font = th_font
        cell.fill = th_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = th_border

    # Data Rows (Row 9 to 38)
    for r_idx, row_data in enumerate(data_rows, start=9):
        ws.row_dimensions[r_idx].height = 34
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = td_font
            cell.border = thin_border
            if c_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 2:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 9: # Status
                cell.font = pass_font
                cell.fill = pass_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Column Widths
    col_widths = {
        1: 8,   # Sr. No.
        2: 12,  # Test Case ID
        3: 26,  # Test Scenario
        4: 24,  # Test Data
        5: 22,  # Preconditions
        6: 28,  # Test Steps
        7: 30,  # Expected Result
        8: 30,  # Actual Result
        9: 10   # Status
    }
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

# 1. Master Complete Workbook (A1_B4_52_PRANAV_RATHI_TESTCASES.xlsx)
wb_master = openpyxl.Workbook()
ws_bb = wb_master.active
build_college_sheet(ws_bb, "Black Box Testing", "Black Box Testing", black_box_rows)

ws_wb = wb_master.create_sheet()
build_college_sheet(ws_wb, "White Box Testing", "White Box Testing", white_box_rows)

out_master = r"c:\Users\prana\OneDrive\Desktop\A1_B4_52_PRANAV_RATHI_TESTCASES.xlsx"
wb_master.save(out_master)
print(f"Saved: {out_master}")

# 2. Individual Black Box Workbook
wb_single_bb = openpyxl.Workbook()
ws1 = wb_single_bb.active
build_college_sheet(ws1, "Black Box Testing", "Black Box Testing", black_box_rows)
out_bb = r"c:\Users\prana\OneDrive\Desktop\A1_B4_52_PRANAV_RATHI_BLACK_BOX.xlsx"
wb_single_bb.save(out_bb)
print(f"Saved: {out_bb}")

# 3. Individual White Box Workbook
wb_single_wb = openpyxl.Workbook()
ws2 = wb_single_wb.active
build_college_sheet(ws2, "White Box Testing", "White Box Testing", white_box_rows)
out_wb = r"c:\Users\prana\OneDrive\Desktop\A1_B4_52_PRANAV_RATHI_WHITE_BOX.xlsx"
wb_single_wb.save(out_wb)
print(f"Saved: {out_wb}")

# 4. Generate matching CSVs
import csv

with open(r"c:\Users\prana\OneDrive\Desktop\A1_B4_52_PRANAV_RATHI_BLACK_BOX.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(headers)
    w.writerows(black_box_rows)

with open(r"c:\Users\prana\OneDrive\Desktop\A1_B4_52_PRANAV_RATHI_WHITE_BOX.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(headers)
    w.writerows(white_box_rows)

print("All workbooks and CSVs generated successfully.")
