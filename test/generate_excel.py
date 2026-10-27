import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_styled_sheet(ws, title, csv_path):
    ws.title = title
    ws.views.sheetView[0].showGridLines = True
    
    # Header styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark blue
    data_font = Font(name="Calibri", size=10)
    pass_font = Font(name="Calibri", size=10, bold=True, color="166534")
    pass_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Light green
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    with open(csv_path, mode='r', encoding='utf-8') as f:
        reader = csv.reader(f)
        rows = list(reader)
        
    for r_idx, row in enumerate(rows, start=1):
        ws.append(row)
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            if r_idx == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.font = data_font
                if str(val).strip().lower() == "pass":
                    cell.font = pass_font
                    cell.fill = pass_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx == 1:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(vertical="center")
                    
    # Auto-fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)
        
    ws.row_dimensions[1].height = 28

# 1. Individual Black Box Excel
wb_bb = openpyxl.Workbook()
ws_bb = wb_bb.active
create_styled_sheet(ws_bb, "Black Box Testing", "test/PBL_Black_Box_Testing.csv")
wb_bb.save("test/PBL_Black_Box_Testing.xlsx")
print("Saved test/PBL_Black_Box_Testing.xlsx")

# 2. Individual White Box Excel
wb_wb = openpyxl.Workbook()
ws_wb = wb_wb.active
create_styled_sheet(ws_wb, "White Box Testing", "test/PBL_White_Box_Testing.csv")
wb_wb.save("test/PBL_White_Box_Testing.xlsx")
print("Saved test/PBL_White_Box_Testing.xlsx")

# 3. Combined Master Workbook
wb_all = openpyxl.Workbook()
ws1 = wb_all.active
create_styled_sheet(ws1, "Black Box Testing (30 Cases)", "test/PBL_Black_Box_Testing.csv")
ws2 = wb_all.create_sheet()
create_styled_sheet(ws2, "White Box Testing (30 Cases)", "test/PBL_White_Box_Testing.csv")
wb_all.save("test/PBL_Testing_Suite_Complete.xlsx")
print("Saved test/PBL_Testing_Suite_Complete.xlsx")
